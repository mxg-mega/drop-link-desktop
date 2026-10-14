"""
LABEL MAKER  (v2 - automatic row height + live preview)
-------------------------------------------------------
Creates a Word document containing a table of labels.

Setup (one time):   pip install python-docx
Run:                double-click this file, or run:  python label_maker.py
"""

import math
import os
import subprocess
import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter import font as tkfont

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Length, Pt

BOLD_CHOICES = ["Same as default", "None", "First line only", "All text"]
DEFAULT_BOLD_CHOICES = ["None", "First line only", "All text"]
FONTS = ["Times New Roman", "Arial", "Calibri", "Cambria", "Verdana", "Courier New"]

# Page = US Letter (Word's default). Change to 21.0 x 29.7 for A4.
PAGE_W_CM, PAGE_H_CM = 21.59, 27.94
MARGIN_CM = 1.0
PAD_CM = 0.1          # space above and below the text inside each cell
LINE_FACTOR = 1.15    # how tall one line of text is, relative to font size
PT_PER_CM = 28.3465


def bold_for(label, default_bold, line_index):
    mode = default_bold if label["bold"] == "Same as default" else label["bold"]
    return mode == "All text" or (mode == "First line only" and line_index == 0)


# ----------------------------------------------------------------------
# Word document builder
# ----------------------------------------------------------------------
def build_document(path, labels, s):
    flat = []
    for lab in labels:
        flat.extend([lab] * lab["qty"])
    cols = s["cols"]
    rows = math.ceil(len(flat) / cols)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
    if s["landscape"]:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(MARGIN_CM))

    page_width = sec.page_width or Cm(PAGE_W_CM)
    left_margin = sec.left_margin or Cm(MARGIN_CM)
    right_margin = sec.right_margin or Cm(MARGIN_CM)
    col_w = int((float(page_width) - float(left_margin) - float(right_margin)) / cols)

    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Small space above/below text in every cell
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "bottom"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(PAD_CM * 567)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    look = table._tbl.tblPr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(mar)
    else:
        table._tbl.tblPr.append(mar)

    for r in table.rows:
        if s["rowh"] > 0:                      # minimum height only
            r.height = Cm(s["rowh"])
            r.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell in r.cells:                   # otherwise row grows with its text
            cell.width = col_w

    for i, lab in enumerate(flat):
        cell = table.cell(i // cols, i % cols)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        size = lab["size"] or s["size"]
        for n, line in enumerate(lab["text"].split("\n")):
            para = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(line)
            run.font.name = s["font"]
            run._element.rPr.rFonts.set(qn("w:eastAsia"), s["font"])
            run.font.size = Pt(size)
            run.bold = bold_for(lab, s["bold"], n)

    doc.save(path)
    return len(flat), rows


# ----------------------------------------------------------------------
# One label entry row in the window
# ----------------------------------------------------------------------
class LabelRow:
    def __init__(self, parent, remove_callback):
        self.frame = ttk.Frame(parent, padding=(0, 3))
        self.text = tk.Text(self.frame, width=20, height=3, font=("Segoe UI", 10))
        self.qty = ttk.Spinbox(self.frame, from_=1, to=2000, width=6)
        self.qty.set("10")
        self.size = ttk.Spinbox(self.frame, from_=4, to=72, width=5)
        self.bold = ttk.Combobox(self.frame, values=BOLD_CHOICES, state="readonly", width=14)
        self.bold.set(BOLD_CHOICES[0])
        self.btn = ttk.Button(self.frame, text="X", width=3,
                              command=lambda: remove_callback(self))
        for col, w in enumerate((self.text, self.qty, self.size, self.bold, self.btn)):
            w.grid(row=0, column=col, padx=3)
        self.frame.pack(fill="x")

    def raw(self):
        return (self.text.get("1.0", "end"), self.qty.get(), self.size.get(), self.bold.get())

    def get(self):
        text = self.text.get("1.0", "end").strip("\n")
        if not text.strip():
            return None
        try:
            qty = int(self.qty.get())
            assert qty >= 1
        except (ValueError, AssertionError):
            raise ValueError("'How many' must be a whole number (1 or more).")
        size_txt = self.size.get().strip()
        try:
            size = float(size_txt) if size_txt else None
        except ValueError:
            raise ValueError("A label's font size must be a number (or empty).")
        return {"text": text, "qty": qty, "size": size, "bold": self.bold.get()}


def enable_wheel(canvas):
    canvas.bind("<Enter>", lambda e: canvas.bind_all(
        "<MouseWheel>", lambda ev: canvas.yview_scroll(int(-ev.delta / 120), "units")))
    canvas.bind("<Leave>", lambda e: canvas.unbind_all("<MouseWheel>"))


# ----------------------------------------------------------------------
# Main window
# ----------------------------------------------------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Label Maker")
        self.geometry("1250x720")
        self.rows = []
        self._fonts = {}
        self._last_sig = None

        left = ttk.Frame(self)
        left.pack(side="left", fill="y", padx=(10, 5), pady=10)
        right = ttk.Frame(self)
        right.pack(side="left", fill="both", expand=True, padx=(5, 10), pady=10)

        # ---- Settings ----
        box = ttk.LabelFrame(left, text="1. Settings for the whole table", padding=8)
        box.pack(fill="x")

        ttk.Label(box, text="Font:").grid(row=0, column=0, sticky="w")
        self.font = ttk.Combobox(box, values=FONTS, width=18)
        self.font.set("Times New Roman")
        self.font.grid(row=0, column=1, padx=6, pady=2)
        ttk.Label(box, text="Font size:").grid(row=0, column=2, sticky="w")
        self.size = ttk.Spinbox(box, from_=4, to=72, width=5)
        self.size.set("11")
        self.size.grid(row=0, column=3, padx=6)

        ttk.Label(box, text="Bold:").grid(row=1, column=0, sticky="w")
        self.bold = ttk.Combobox(box, values=DEFAULT_BOLD_CHOICES, state="readonly", width=16)
        self.bold.set("None")
        self.bold.grid(row=1, column=1, padx=6, pady=2, sticky="w")
        ttk.Label(box, text="Columns:").grid(row=1, column=2, sticky="w")
        self.cols = ttk.Spinbox(box, from_=1, to=30, width=5)
        self.cols.set("10")
        self.cols.grid(row=1, column=3, padx=6)

        ttk.Label(box, text="Min row height (cm):").grid(row=2, column=0, sticky="w")
        self.rowh = ttk.Entry(box, width=8)
        self.rowh.grid(row=2, column=1, padx=6, pady=2, sticky="w")
        ttk.Label(box, text="empty = automatic", foreground="gray")\
            .grid(row=2, column=2, columnspan=2, sticky="w")

        self.landscape = tk.BooleanVar(value=True)
        ttk.Checkbutton(box, text="Landscape page", variable=self.landscape)\
            .grid(row=3, column=0, columnspan=2, sticky="w", pady=(4, 0))

        # ---- Labels ----
        ttk.Label(left, text="2. Your labels  (Enter = new line inside a label)",
                  font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(10, 0))
        head = ttk.Frame(left)
        head.pack(fill="x")
        for col, (txt, w) in enumerate([("Label text", 22), ("How many", 7),
                                        ("Size*", 6), ("Bold", 15), ("", 3)]):
            ttk.Label(head, text=txt, width=w).grid(row=0, column=col, padx=3)

        holder = ttk.Frame(left)
        holder.pack(fill="both", expand=True)
        canvas = tk.Canvas(holder, highlightthickness=0, width=560)
        scroll = ttk.Scrollbar(holder, orient="vertical", command=canvas.yview)
        self.inner = ttk.Frame(canvas)
        self.inner.bind("<Configure>",
                        lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.inner, anchor="nw")
        canvas.configure(yscrollcommand=scroll.set)
        canvas.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        enable_wheel(canvas)

        ttk.Label(left, text="*Leave size empty to use the default above.",
                  foreground="gray").pack(anchor="w")
        bar = ttk.Frame(left, padding=6)
        bar.pack(fill="x")
        ttk.Button(bar, text="+ Add another label", command=self.add_row).pack(side="left")
        ttk.Button(bar, text="Create Word document", command=self.generate).pack(side="right")

        # ---- Preview ----
        ttk.Label(right, text="Live preview (updates as you type)",
                  font=("Segoe UI", 9, "bold")).pack(anchor="w")
        self.status = ttk.Label(right, text="", foreground="gray")
        self.status.pack(anchor="w")
        pframe = ttk.Frame(right)
        pframe.pack(fill="both", expand=True)
        self.preview = tk.Canvas(pframe, bg="#8a8a8a", highlightthickness=0)
        pscroll = ttk.Scrollbar(pframe, orient="vertical", command=self.preview.yview)
        self.preview.configure(yscrollcommand=pscroll.set)
        self.preview.pack(side="left", fill="both", expand=True)
        pscroll.pack(side="right", fill="y")
        enable_wheel(self.preview)

        self.add_row()
        self.rows[0].text.insert("1.0", "M\n5/10/26")
        self.after(300, self.poll)

    # ---------- rows ----------
    def add_row(self):
        self.rows.append(LabelRow(self.inner, self.remove_row))

    def remove_row(self, row):
        row.frame.destroy()
        self.rows.remove(row)

    # ---------- reading inputs ----------
    def read_settings(self):
        try:
            cols = int(self.cols.get())
            assert 1 <= cols <= 30
            size = float(self.size.get())
            rowh_txt = self.rowh.get().strip()
            rowh = float(rowh_txt) if rowh_txt else 0.0
        except (ValueError, AssertionError):
            raise ValueError("Please check the numbers in the settings "
                             "(columns 1-30, font size, row height).")
        return {"cols": cols, "size": size, "rowh": rowh, "font": self.font.get(),
                "bold": self.bold.get(), "landscape": self.landscape.get()}

    def read_labels(self):
        labels = [l for l in (r.get() for r in self.rows) if l]
        return labels

    # ---------- live preview ----------
    def poll(self):
        sig = (self.font.get(), self.size.get(), self.bold.get(), self.cols.get(),
               self.rowh.get(), self.landscape.get(), self.preview.winfo_width(),
               tuple(r.raw() for r in self.rows))
        if sig != self._last_sig:
            self._last_sig = sig
            try:
                self.draw_preview(self.read_labels(), self.read_settings())
            except ValueError as e:
                self.preview.delete("all")
                self.preview.create_text(20, 20, anchor="nw", fill="white", text=str(e))
                self.status.config(text="")
        self.after(300, self.poll)

    def get_font(self, family, px, bold):
        key = (family, px, bold)
        if key not in self._fonts:
            self._fonts[key] = tkfont.Font(family=family, size=-px,
                                           weight="bold" if bold else "normal")
        return self._fonts[key]

    def draw_preview(self, labels, s):
        c = self.preview
        c.delete("all")
        flat = []
        for lab in labels:
            flat.extend([lab] * lab["qty"])
        if not flat:
            self.status.config(text="Type a label to see the preview.")
            c.configure(scrollregion=(0, 0, 1, 1))
            return

        pw_cm, ph_cm = (PAGE_H_CM, PAGE_W_CM) if s["landscape"] else (PAGE_W_CM, PAGE_H_CM)
        margin_px = 12
        W = max(c.winfo_width(), 300)
        k = (W - 2 * margin_px) / pw_cm           # pixels per cm
        px_per_pt = k / PT_PER_CM
        cols = s["cols"]
        col_w = (pw_cm - 2 * MARGIN_CM) / cols * k
        usable_h_pt = (ph_cm - 2 * MARGIN_CM) * PT_PER_CM
        pad_pt = PAD_CM * PT_PER_CM

        def page_top(p):
            return margin_px + p * (ph_cm * k + margin_px)

        def new_page(p):
            t = page_top(p)
            c.create_rectangle(margin_px, t, margin_px + pw_cm * k, t + ph_cm * k,
                               fill="white", outline="#555")

        page, y_pt = 0, 0.0
        new_page(0)
        nrows = math.ceil(len(flat) / cols)
        for r in range(nrows):
            cells = flat[r * cols:(r + 1) * cols]
            # Row height grows with its tallest cell (like Word)
            h_pt = s["rowh"] * PT_PER_CM
            for lab in cells:
                size = lab["size"] or s["size"]
                n = len(lab["text"].split("\n"))
                h_pt = max(h_pt, n * size * LINE_FACTOR + 2 * pad_pt)
            if len(cells) < cols:
                h_pt = max(h_pt, s["size"] * LINE_FACTOR + 2 * pad_pt)

            if y_pt + h_pt > usable_h_pt and y_pt > 0:
                page += 1
                y_pt = 0.0
                new_page(page)

            top = page_top(page) + MARGIN_CM * k + y_pt * px_per_pt
            bottom = top + h_pt * px_per_pt
            x0 = margin_px + MARGIN_CM * k
            for ci in range(cols):
                c.create_rectangle(x0 + ci * col_w, top, x0 + (ci + 1) * col_w, bottom,
                                   outline="#333")
                if ci < len(cells):
                    lab = cells[ci]
                    size = lab["size"] or s["size"]
                    lines = lab["text"].split("\n")
                    lh = size * LINE_FACTOR * px_per_pt
                    cy = (top + bottom) / 2 - lh * len(lines) / 2 + lh / 2
                    px = max(1, int(round(size * px_per_pt)))
                    for n, line in enumerate(lines):
                        f = self.get_font(s["font"], px, bold_for(lab, s["bold"], n))
                        c.create_text(x0 + (ci + 0.5) * col_w, cy + n * lh,
                                      text=line, font=f, fill="black")
            y_pt += h_pt

        c.configure(scrollregion=(0, 0, W, page_top(page + 1)))
        self.status.config(text=f"{len(flat)} labels  |  {nrows} rows  |  "
                                f"{page + 1} page(s)   (preview is approximate)")

    # ---------- create document ----------
    def generate(self):
        try:
            labels = self.read_labels()
            s = self.read_settings()
        except ValueError as e:
            messagebox.showerror("Label Maker", str(e))
            return
        if not labels:
            messagebox.showwarning("Label Maker", "Please type at least one label.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".docx", filetypes=[("Word document", "*.docx")],
            initialfile="labels.docx", title="Save labels as")
        if not path:
            return
        try:
            total, rows = build_document(path, labels, s)
        except PermissionError:
            messagebox.showerror("Label Maker",
                                 "Couldn't save. If that file is open in Word, close it and try again.")
            return

        messagebox.showinfo("Label Maker", f"Done! {total} labels in {rows} rows.")
        try:
            if sys.platform.startswith("win"):
                startfile = getattr(os, "startfile", None)
                if startfile is not None:
                    startfile(path)
            elif sys.platform == "darwin":
                subprocess.run(["open", path], check=False)
            else:
                subprocess.run(["xdg-open", path], check=False)
        except Exception:
            pass


if __name__ == "__main__":
    App().mainloop()